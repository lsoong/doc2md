"""테스트용 샘플 문서 생성 (docx/pptx/xlsx/pdf).

실행: python tests/make_samples.py [출력디렉터리]
표·제목·목록이 섞인 한국어 문서를 만들어 엔진별 변환 품질을 눈으로 비교할 수 있게 한다.
"""

from __future__ import annotations

import sys
from pathlib import Path

TITLE = "2026년 3분기 문서 변환 파일럿 보고서"
SECTIONS = [
    ("배경", "사내 문서를 RAG 인덱스에 넣기 전에 Markdown 으로 표준화할 필요가 있다."),
    ("범위", "PDF, Word, PowerPoint, Excel 네 가지 포맷을 대상으로 한다."),
]
TABLE = [
    ["엔진", "포맷", "평균 처리시간(초)", "표 정합"],
    ["MinerU", "PDF", "12.4", "0.98"],
    ["Docling", "PDF/Office", "8.1", "0.96"],
    ["Marker", "PDF", "6.7", "0.91"],
    ["MarkItDown", "전포맷", "0.4", "0.62"],
]
BULLETS = ["정확도 우선: MinerU", "범용성 우선: Docling", "속도 우선: PyMuPDF4LLM"]


def make_docx(path: Path) -> None:
    from docx import Document

    doc = Document()
    doc.add_heading(TITLE, level=1)
    for heading, body in SECTIONS:
        doc.add_heading(heading, level=2)
        doc.add_paragraph(body)
    doc.add_heading("엔진 비교", level=2)
    table = doc.add_table(rows=len(TABLE), cols=len(TABLE[0]))
    table.style = "Table Grid"
    for r, row in enumerate(TABLE):
        for c, cell in enumerate(row):
            table.cell(r, c).text = cell
    doc.add_heading("결론", level=2)
    for bullet in BULLETS:
        doc.add_paragraph(bullet, style="List Bullet")
    doc.save(path)


def make_pptx(path: Path) -> None:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = TITLE
    slide.placeholders[1].text = "장인 / SYAI"

    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "결론"
    body = slide.placeholders[1].text_frame
    body.text = BULLETS[0]
    for bullet in BULLETS[1:]:
        body.add_paragraph().text = bullet

    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "엔진 비교"
    shape = slide.shapes.add_table(
        len(TABLE), len(TABLE[0]), Inches(0.5), Inches(1.8), Inches(9), Inches(3)
    )
    for r, row in enumerate(TABLE):
        for c, cell in enumerate(row):
            shape.table.cell(r, c).text = cell
    prs.save(path)


def make_xlsx(path: Path) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "엔진비교"
    for row in TABLE:
        ws.append(row)
    ws2 = wb.create_sheet("결론")
    ws2.append(["순위", "권고"])
    for index, bullet in enumerate(BULLETS, start=1):
        ws2.append([index, bullet])
    wb.save(path)


def make_pdf(path: Path) -> None:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    font = "Helvetica"
    for candidate in (
        "/usr/share/fonts/truetype/nanum/NanumGothic.ttf",
        "/mnt/c/Windows/Fonts/malgun.ttf",
        "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
    ):
        if Path(candidate).exists():
            try:
                pdfmetrics.registerFont(TTFont("KR", candidate))
                font = "KR"
                break
            except Exception:  # noqa: BLE001 - 폰트 등록 실패는 무시하고 기본 폰트로
                continue

    styles = getSampleStyleSheet()
    h1 = ParagraphStyle("h1", parent=styles["Heading1"], fontName=font, fontSize=18)
    h2 = ParagraphStyle("h2", parent=styles["Heading2"], fontName=font, fontSize=14)
    body = ParagraphStyle("body", parent=styles["BodyText"], fontName=font, fontSize=10.5)

    story = [Paragraph(TITLE, h1), Spacer(1, 10)]
    for heading, text in SECTIONS:
        story += [Paragraph(heading, h2), Paragraph(text, body), Spacer(1, 6)]
    story.append(Paragraph("엔진 비교", h2))
    table = Table(TABLE, hAlign="LEFT")
    table.setStyle(
        TableStyle(
            [
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                ("BACKGROUND", (0, 0), (-1, 0), colors.whitesmoke),
                ("FONTNAME", (0, 0), (-1, -1), font),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
            ]
        )
    )
    story += [table, Spacer(1, 10), Paragraph("결론", h2)]
    for bullet in BULLETS:
        story.append(Paragraph(f"• {bullet}", body))
    SimpleDocTemplate(str(path), pagesize=A4, title=TITLE).build(story)


def main() -> None:
    out = Path(sys.argv[1] if len(sys.argv) > 1 else "tests/samples")
    out.mkdir(parents=True, exist_ok=True)
    make_docx(out / "sample.docx")
    make_pptx(out / "sample.pptx")
    make_xlsx(out / "sample.xlsx")
    make_pdf(out / "sample.pdf")
    for path in sorted(out.iterdir()):
        print(f"{path}  {path.stat().st_size:,} bytes")


if __name__ == "__main__":
    main()
